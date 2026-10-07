"""
ARBOSSTENT BOT — ichki hisob-kitob boti
Kirim/chiqim, asbob ijarasi, qarzlar, hisobotlar, Excel va avtomatik zaxira.
Sozlamalar muhit o'zgaruvchilari orqali: BOT_TOKEN, ALLOWED_IDS, DB_PATH
"""
import asyncio
import io
import logging
import os
import sqlite3
from contextlib import contextmanager
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

from aiogram import Bot, Dispatcher, F
from aiogram.filters import CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (
    BufferedInputFile,
    CallbackQuery,
    FSInputFile,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    KeyboardButton,
    Message,
    ReplyKeyboardMarkup,
)
from openpyxl import Workbook
from openpyxl.styles import Font

# ====================== SOZLAMALAR ======================
BOT_TOKEN = os.getenv("BOT_TOKEN", "")
ALLOWED = {int(x) for x in os.getenv("ALLOWED_IDS", "").replace(" ", "").split(",") if x}
DB = os.getenv("DB_PATH", "arbosstent.db")
TZ = ZoneInfo("Asia/Tashkent")

# Avtomatik xabarlar (Toshkent vaqti bilan, soat)
HOUR_REMINDER = 9    # ertalab: qaytarilishi kerak bo'lgan asboblar
HOUR_REPORT = 21     # kechqurun: kunlik hisobot
HOUR_BACKUP = 22     # kechasi: baza nusxasi

SERVICES = ["Asbob ijarasi", "Musor olib ketish", "Kran", "Gruzchik", "Asbob sotuvi", "Boshqa"]
EXPENSES = ["Benzin/yoqilg'i", "Maosh", "Ta'mir", "Asbob xaridi", "Reklama", "Boshqa"]
PAYS = ["Naqd", "Karta", "O'tkazma", "Qarz"]
SOURCES = ["Onlayn", "Oflayn"]

B_IN, B_OUT = "➕ Kirim", "➖ Chiqim"
B_RENT, B_ACTIVE = "🛠 Ijara berish", "📦 Ijaradagilar"
B_DEBT, B_REP = "💳 Qarzlar", "📊 Hisobot"
B_LAST, B_XLS = "📋 Oxirgilar", "📁 Excel"
B_BACKUP, B_DEL = "🛟 Zaxira", "🗑 Oxirgisini o'chirish"
B_CANCEL = "❌ Bekor qilish"
PERIODS = ["Bugun", "Kecha", "7 kun", "Shu oy", "O'tgan oy"]

logging.basicConfig(level=logging.INFO)


# ====================== YORDAMCHI ======================
def now() -> datetime:
    return datetime.now(TZ)


def fmt(n) -> str:
    return f"{int(n or 0):,}".replace(",", " ")


def parse_int(text: str):
    digits = "".join(ch for ch in text if ch.isdigit())
    if not digits or int(digits) <= 0:
        return None
    return int(digits)


def kb(items, per_row=2, extra=None) -> ReplyKeyboardMarkup:
    rows, row = [], []
    for it in items:
        row.append(KeyboardButton(text=it))
        if len(row) == per_row:
            rows.append(row)
            row = []
    if row:
        rows.append(row)
    if extra:
        rows.append([KeyboardButton(text=extra)])
    return ReplyKeyboardMarkup(keyboard=rows, resize_keyboard=True)


MAIN = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text=B_IN), KeyboardButton(text=B_OUT)],
        [KeyboardButton(text=B_RENT), KeyboardButton(text=B_ACTIVE)],
        [KeyboardButton(text=B_DEBT), KeyboardButton(text=B_REP)],
        [KeyboardButton(text=B_LAST), KeyboardButton(text=B_XLS)],
        [KeyboardButton(text=B_BACKUP), KeyboardButton(text=B_DEL)],
    ],
    resize_keyboard=True,
)
CANCEL_KB = kb([], extra=B_CANCEL)


# ====================== BAZA ======================
@contextmanager
def conn():
    con = sqlite3.connect(DB)
    try:
        yield con
        con.commit()
    finally:
        con.close()


def init_db():
    folder = os.path.dirname(DB)
    if folder:
        os.makedirs(folder, exist_ok=True)
    with conn() as con:
        con.execute(
            """CREATE TABLE IF NOT EXISTS records(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                date TEXT, time TEXT, kind TEXT, category TEXT, amount INTEGER,
                pay TEXT, source TEXT, client TEXT, note TEXT,
                paid INTEGER DEFAULT 1, paid_date TEXT, rental_id INTEGER, user_id INTEGER)"""
        )
        con.execute(
            """CREATE TABLE IF NOT EXISTS rentals(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                tool TEXT, client TEXT, phone TEXT, start_date TEXT, due_date TEXT,
                days INTEGER, price_day INTEGER, total INTEGER, deposit INTEGER,
                returned INTEGER DEFAULT 0, returned_date TEXT, user_id INTEGER)"""
        )


def add_record(kind, category, amount, pay, source, client, note, user_id, rental_id=None):
    n = now()
    paid = 0 if pay == "Qarz" else 1
    with conn() as con:
        cur = con.execute(
            """INSERT INTO records(date,time,kind,category,amount,pay,source,client,note,paid,rental_id,user_id)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
            (n.strftime("%Y-%m-%d"), n.strftime("%H:%M"), kind, category, amount,
             pay, source, client, note, paid, rental_id, user_id),
        )
        return cur.lastrowid


def period_range(name: str):
    today = now().date()
    if name == "Bugun":
        return today, today
    if name == "Kecha":
        d = today - timedelta(days=1)
        return d, d
    if name == "7 kun":
        return today - timedelta(days=6), today
    if name == "Shu oy":
        return today.replace(day=1), today
    last_end = today.replace(day=1) - timedelta(days=1)  # O'tgan oy
    return last_end.replace(day=1), last_end


def build_report(start: str, end: str, title: str) -> str:
    q = "FROM records WHERE kind=? AND date BETWEEN ? AND ?"
    with conn() as con:
        inc = con.execute(f"SELECT category, SUM(amount), COUNT(*) {q} GROUP BY category ORDER BY 2 DESC",
                          ("kirim", start, end)).fetchall()
        exp = con.execute(f"SELECT category, SUM(amount), COUNT(*) {q} GROUP BY category ORDER BY 2 DESC",
                          ("chiqim", start, end)).fetchall()
        pays = con.execute(f"SELECT pay, SUM(amount) {q} GROUP BY pay", ("kirim", start, end)).fetchall()
        srcs = con.execute(f"SELECT source, SUM(amount) {q} GROUP BY source", ("kirim", start, end)).fetchall()
        debt = con.execute(
            f"SELECT COALESCE(SUM(amount),0) {q} AND pay='Qarz' AND paid=0", ("kirim", start, end)
        ).fetchone()[0]
        all_debt = con.execute(
            "SELECT COALESCE(SUM(amount),0) FROM records WHERE kind='kirim' AND pay='Qarz' AND paid=0"
        ).fetchone()[0]
        out_now = con.execute("SELECT COUNT(*) FROM rentals WHERE returned=0").fetchone()[0]
        overdue = con.execute(
            "SELECT COUNT(*) FROM rentals WHERE returned=0 AND due_date < ?", (now().strftime("%Y-%m-%d"),)
        ).fetchone()[0]

    t_in = sum(r[1] for r in inc)
    t_out = sum(r[1] for r in exp)
    L = [f"📊 <b>{title}</b>", ""]
    L.append("<b>Kirim:</b>")
    L += [f"  • {c}: {fmt(a)} ({n} ta)" for c, a, n in inc] or ["  —"]
    L.append(f"  <b>Jami kirim: {fmt(t_in)}</b>")
    if pays:
        L.append("  " + " | ".join(f"{p}: {fmt(a)}" for p, a in pays))
    if srcs:
        L.append("  " + " | ".join(f"{s or '-'}: {fmt(a)}" for s, a in srcs))
    L.append("")
    L.append("<b>Chiqim:</b>")
    L += [f"  • {c}: {fmt(a)} ({n} ta)" for c, a, n in exp] or ["  —"]
    L.append(f"  <b>Jami chiqim: {fmt(t_out)}</b>")
    L.append("")
    L.append(f"💰 <b>Foyda: {fmt(t_in - t_out)}</b>")
    if debt:
        L.append(f"💵 Shundan qo'lga tushgan: {fmt(t_in - debt - t_out)}")
        L.append(f"⚠️ Davrdagi to'lanmagan qarz: {fmt(debt)}")
    L.append("")
    L.append(f"💳 Jami qarzdorlik (hozir): {fmt(all_debt)}")
    L.append(f"📦 Ijarada: {out_now} ta" + (f" (⚠️ muddati o'tgan: {overdue})" if overdue else ""))
    return "\n".join(L)


# ====================== EXCEL ======================
def build_excel() -> bytes:
    wb = Workbook()
    bold = Font(bold=True)

    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font = bold
        for r in rows:
            ws.append(list(r))
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w

    with conn() as con:
        recs = con.execute(
            """SELECT date,time,kind,category,amount,pay,source,client,note,
               CASE WHEN pay='Qarz' AND paid=0 THEN 'to''lanmagan' ELSE '' END
               FROM records ORDER BY id"""
        ).fetchall()
        rents = con.execute(
            """SELECT id,tool,client,phone,start_date,due_date,days,price_day,total,deposit,
               CASE WHEN returned=1 THEN 'qaytarilgan' ELSE 'ijarada' END, returned_date
               FROM rentals ORDER BY id"""
        ).fetchall()
        debts = con.execute(
            "SELECT date,client,category,amount,note FROM records WHERE kind='kirim' AND pay='Qarz' AND paid=0 ORDER BY date"
        ).fetchall()

    ws = wb.active
    ws.title = "Kirim-Chiqim"
    sheet(ws, ["Sana", "Vaqt", "Turi", "Kategoriya", "Summa", "To'lov", "Manba", "Mijoz", "Izoh", "Qarz holati"],
          recs, [12, 8, 9, 20, 14, 11, 10, 20, 30, 12])
    sheet(wb.create_sheet("Ijaralar"),
          ["№", "Asbob", "Mijoz", "Telefon", "Boshlandi", "Qaytarish", "Kun", "Kunlik narx", "Jami", "Garov", "Holat", "Qaytgan sana"],
          rents, [5, 22, 20, 16, 12, 12, 6, 12, 12, 12, 12, 12])
    sheet(wb.create_sheet("Qarzlar"), ["Sana", "Mijoz", "Xizmat", "Summa", "Izoh"], debts, [12, 22, 20, 14, 30])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_backup() -> str:
    path = DB + ".bak"
    src, dst = sqlite3.connect(DB), sqlite3.connect(path)
    try:
        src.backup(dst)
    finally:
        dst.close()
        src.close()
    return path


# ====================== HOLATLAR ======================
class Rec(StatesGroup):
    category = State()
    amount = State()
    pay = State()
    source = State()
    client = State()
    note = State()


class Rent(StatesGroup):
    tool = State()
    client = State()
    phone = State()
    days = State()
    price = State()
    deposit = State()
    pay = State()


dp = Dispatcher(storage=MemoryStorage())


# ====================== RUXSAT ======================
async def access(handler, event, data):
    uid = event.from_user.id
    if uid not in ALLOWED:
        text = f"⛔ Ruxsat yo'q.\nSizning ID: {uid}\nShu raqamni bot egasiga yuboring."
        if isinstance(event, CallbackQuery):
            await event.answer(text, show_alert=True)
        else:
            await event.answer(text)
        return
    return await handler(event, data)


dp.message.outer_middleware(access)
dp.callback_query.outer_middleware(access)


# ====================== ASOSIY ======================
@dp.message(CommandStart())
async def start(m: Message, state: FSMContext):
    await state.clear()
    await m.answer("Assalomu alaykum! 📒 ARBOSSTENT hisob boti.\nMenyudan tanlang 👇", reply_markup=MAIN)


@dp.message(F.text == B_CANCEL)
async def cancel(m: Message, state: FSMContext):
    await state.clear()
    await m.answer("Bekor qilindi.", reply_markup=MAIN)


# ---------- kirim / chiqim ----------
@dp.message(F.text == B_IN)
async def new_in(m: Message, state: FSMContext):
    await state.clear()
    await state.update_data(kind="kirim")
    await state.set_state(Rec.category)
    await m.answer("Qaysi xizmat/mahsulot?", reply_markup=kb(SERVICES, 2, B_CANCEL))


@dp.message(F.text == B_OUT)
async def new_out(m: Message, state: FSMContext):
    await state.clear()
    await state.update_data(kind="chiqim")
    await state.set_state(Rec.category)
    await m.answer("Chiqim turi?", reply_markup=kb(EXPENSES, 2, B_CANCEL))


# ---------- boshqa menyu tugmalari ----------
@dp.message(F.text == B_REP)
async def rep_menu(m: Message, state: FSMContext):
    await state.clear()
    await m.answer("Qaysi davr?", reply_markup=kb(PERIODS, 3, B_CANCEL))


@dp.message(F.text.in_(PERIODS))
async def rep_show(m: Message):
    s, e = period_range(m.text)
    title = f"{m.text}: {s.strftime('%d.%m.%Y')}" + ("" if s == e else f" — {e.strftime('%d.%m.%Y')}")
    await m.answer(build_report(str(s), str(e), title), parse_mode="HTML", reply_markup=MAIN)


@dp.message(F.text == B_LAST)
async def last(m: Message, state: FSMContext):
    await state.clear()
    with conn() as con:
        rows = con.execute(
            "SELECT date,time,kind,category,amount,pay,client FROM records ORDER BY id DESC LIMIT 15"
        ).fetchall()
    if not rows:
        await m.answer("Hali yozuv yo'q.")
        return
    out = []
    for d, t, k, c, a, p, cl in rows:
        out.append(f"{d[5:]} {t} {'➕' if k == 'kirim' else '➖'} {c}: {fmt(a)} ({p}) {cl or ''}".strip())
    await m.answer("\n".join(out))


@dp.message(F.text == B_DEL)
async def del_ask(m: Message, state: FSMContext):
    await state.clear()
    with conn() as con:
        r = con.execute("SELECT id,date,category,amount FROM records ORDER BY id DESC LIMIT 1").fetchone()
    if not r:
        await m.answer("O'chirish uchun yozuv yo'q.")
        return
    markup = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="🗑 Ha, o'chir", callback_data=f"del:{r[0]}"),
        InlineKeyboardButton(text="Yo'q", callback_data="noop"),
    ]])
    await m.answer(f"Oxirgi yozuv o'chirilsinmi?\n{r[1]} — {r[2]}: {fmt(r[3])}", reply_markup=markup)


@dp.callback_query(F.data.startswith("del:"))
async def del_do(c: CallbackQuery):
    rid = int(c.data.split(":")[1])
    with conn() as con:
        con.execute("DELETE FROM records WHERE id=?", (rid,))
    await c.message.edit_text("🗑 O'chirildi.")
    await c.answer()


@dp.callback_query(F.data == "noop")
async def noop(c: CallbackQuery):
    await c.message.edit_text("Bekor qilindi.")
    await c.answer()


@dp.message(F.text == B_XLS)
async def excel(m: Message, state: FSMContext):
    await state.clear()
    data = build_excel()
    name = f"arbosstent_{now().strftime('%Y-%m-%d')}.xlsx"
    await m.answer_document(BufferedInputFile(data, filename=name), caption="📁 To'liq hisobot (3 varaq)")


@dp.message(F.text == B_BACKUP)
async def backup_now(m: Message, state: FSMContext):
    await state.clear()
    await m.answer_document(FSInputFile(make_backup(), filename=f"zaxira_{now().strftime('%Y-%m-%d')}.db"),
                            caption="🛟 Baza nusxasi. Saqlab qo'ying.")


# ---------- qarzlar ----------
@dp.message(F.text == B_DEBT)
async def debts(m: Message, state: FSMContext):
    await state.clear()
    with conn() as con:
        rows = con.execute(
            "SELECT id,date,client,category,amount,note FROM records WHERE kind='kirim' AND pay='Qarz' AND paid=0 ORDER BY date LIMIT 20"
        ).fetchall()
        total = con.execute(
            "SELECT COALESCE(SUM(amount),0) FROM records WHERE kind='kirim' AND pay='Qarz' AND paid=0"
        ).fetchone()[0]
    if not rows:
        await m.answer("✅ Qarzdorlar yo'q.")
        return
    await m.answer(f"💳 <b>Jami qarzdorlik: {fmt(total)}</b>", parse_mode="HTML")
    for rid, d, cl, cat, a, note in rows:
        markup = InlineKeyboardMarkup(inline_keyboard=[[
            InlineKeyboardButton(text="✅ To'landi", callback_data=f"paid:{rid}")
        ]])
        await m.answer(f"👤 {cl}\n{d} • {cat}\n💰 {fmt(a)} so'm {('— ' + note) if note else ''}", reply_markup=markup)


@dp.callback_query(F.data.startswith("paid:"))
async def debt_paid(c: CallbackQuery):
    rid = int(c.data.split(":")[1])
    with conn() as con:
        con.execute("UPDATE records SET paid=1, paid_date=? WHERE id=?", (now().strftime("%Y-%m-%d"), rid))
    await c.message.edit_text((c.message.text or "") + "\n\n✅ To'landi")
    await c.answer("Qarz to'langan deb belgilandi")


# ---------- ijaradagilar ----------
def rental_text(r) -> str:
    rid, tool, client, phone, start, due, days, price, total, deposit = r
    today = now().strftime("%Y-%m-%d")
    flag = "⚠️ MUDDATI O'TGAN" if due < today else ("⏰ Bugun qaytadi" if due == today else "")
    return (f"🛠 <b>{tool}</b> #{rid}\n👤 {client} • {phone}\n"
            f"📅 {start} → {due} ({days} kun)\n💰 {fmt(total)} • garov: {fmt(deposit)}\n{flag}").strip()


@dp.message(F.text == B_ACTIVE)
async def active(m: Message, state: FSMContext):
    await state.clear()
    with conn() as con:
        rows = con.execute(
            "SELECT id,tool,client,phone,start_date,due_date,days,price_day,total,deposit FROM rentals WHERE returned=0 ORDER BY due_date LIMIT 25"
        ).fetchall()
    if not rows:
        await m.answer("Hozir ijarada asbob yo'q.")
        return
    for r in rows:
        markup = InlineKeyboardMarkup(inline_keyboard=[[
            InlineKeyboardButton(text="✅ Qaytarildi", callback_data=f"ret:{r[0]}")
        ]])
        await m.answer(rental_text(r), parse_mode="HTML", reply_markup=markup)


@dp.callback_query(F.data.startswith("ret:"))
async def returned(c: CallbackQuery):
    rid = int(c.data.split(":")[1])
    with conn() as con:
        con.execute("UPDATE rentals SET returned=1, returned_date=? WHERE id=?", (now().strftime("%Y-%m-%d"), rid))
    await c.message.edit_text((c.message.text or "") + "\n\n✅ Qaytarildi", reply_markup=None)
    await c.answer("Asbob qaytarildi deb belgilandi")


# ====================== KIRIM/CHIQIM OQIMI ======================
@dp.message(Rec.category, F.text)
async def rec_cat(m: Message, state: FSMContext):
    await state.update_data(category=m.text)
    await state.set_state(Rec.amount)
    await m.answer("Summani yozing (so'm):", reply_markup=CANCEL_KB)


@dp.message(Rec.amount, F.text)
async def rec_amount(m: Message, state: FSMContext):
    amount = parse_int(m.text)
    if not amount:
        await m.answer("Faqat raqam yozing, masalan: 250000")
        return
    await state.update_data(amount=amount)
    d = await state.get_data()
    if d["kind"] == "kirim":
        await state.set_state(Rec.pay)
        await m.answer("To'lov turi?", reply_markup=kb(PAYS, 2, B_CANCEL))
    else:
        await state.update_data(pay="Naqd", source="-", client="")
        await state.set_state(Rec.note)
        await m.answer("Izoh (yoki «-» yuboring):", reply_markup=CANCEL_KB)


@dp.message(Rec.pay, F.text.in_(PAYS))
async def rec_pay(m: Message, state: FSMContext):
    await state.update_data(pay=m.text)
    await state.set_state(Rec.source)
    await m.answer("Buyurtma qayerdan keldi?", reply_markup=kb(SOURCES, 2, B_CANCEL))


@dp.message(Rec.source, F.text.in_(SOURCES))
async def rec_source(m: Message, state: FSMContext):
    await state.update_data(source=m.text)
    await state.set_state(Rec.client)
    await m.answer("Mijoz ismi (yoki «-» yuboring):", reply_markup=CANCEL_KB)


@dp.message(Rec.client, F.text)
async def rec_client(m: Message, state: FSMContext):
    d = await state.get_data()
    if d["pay"] == "Qarz" and m.text.strip() == "-":
        await m.answer("Qarz uchun mijoz ismi shart. Ismini yozing:")
        return
    await state.update_data(client="" if m.text.strip() == "-" else m.text.strip())
    await state.set_state(Rec.note)
    await m.answer("Izoh (yoki «-» yuboring):", reply_markup=CANCEL_KB)


@dp.message(Rec.note, F.text)
async def rec_note(m: Message, state: FSMContext):
    d = await state.get_data()
    await state.clear()
    note = "" if m.text.strip() == "-" else m.text.strip()
    add_record(d["kind"], d["category"], d["amount"], d["pay"], d["source"], d["client"], note, m.from_user.id)
    sign = "➕" if d["kind"] == "kirim" else "➖"
    extra = " ⚠️ qarz sifatida yozildi" if d["pay"] == "Qarz" else ""
    await m.answer(f"✅ Saqlandi: {sign} {d['category']} — {fmt(d['amount'])} so'm ({d['pay']}){extra}",
                   reply_markup=MAIN)


# ====================== IJARA OQIMI ======================
@dp.message(F.text == B_RENT)
async def rent_start(m: Message, state: FSMContext):
    await state.clear()
    await state.set_state(Rent.tool)
    await m.answer("Qaysi asbob? (masalan: perforator)", reply_markup=CANCEL_KB)


@dp.message(Rent.tool, F.text)
async def rent_tool(m: Message, state: FSMContext):
    await state.update_data(tool=m.text.strip())
    await state.set_state(Rent.client)
    await m.answer("Mijoz ismi:")


@dp.message(Rent.client, F.text)
async def rent_client(m: Message, state: FSMContext):
    await state.update_data(client=m.text.strip())
    await state.set_state(Rent.phone)
    await m.answer("Mijoz telefon raqami:")


@dp.message(Rent.phone, F.text)
async def rent_phone(m: Message, state: FSMContext):
    await state.update_data(phone=m.text.strip())
    await state.set_state(Rent.days)
    await m.answer("Necha kunga?")


@dp.message(Rent.days, F.text)
async def rent_days(m: Message, state: FSMContext):
    days = parse_int(m.text)
    if not days or days > 365:
        await m.answer("Kun sonini raqam bilan yozing, masalan: 3")
        return
    await state.update_data(days=days)
    await state.set_state(Rent.price)
    await m.answer("Bir kunlik narx (so'm):")


@dp.message(Rent.price, F.text)
async def rent_price(m: Message, state: FSMContext):
    price = parse_int(m.text)
    if not price:
        await m.answer("Narxni raqam bilan yozing, masalan: 50000")
        return
    await state.update_data(price=price)
    await state.set_state(Rent.deposit)
    await m.answer("Garov (zalog) summasi? Yo'q bo'lsa 0 yozing:")


@dp.message(Rent.deposit, F.text)
async def rent_deposit(m: Message, state: FSMContext):
    digits = "".join(ch for ch in m.text if ch.isdigit())
    await state.update_data(deposit=int(digits) if digits else 0)
    await state.set_state(Rent.pay)
    await m.answer("To'lov turi?", reply_markup=kb(PAYS, 2, B_CANCEL))


@dp.message(Rent.pay, F.text.in_(PAYS))
async def rent_pay(m: Message, state: FSMContext):
    d = await state.get_data()
    await state.clear()
    n = now()
    total = d["days"] * d["price"]
    due = (n + timedelta(days=d["days"])).strftime("%Y-%m-%d")
    with conn() as con:
        cur = con.execute(
            """INSERT INTO rentals(tool,client,phone,start_date,due_date,days,price_day,total,deposit,user_id)
               VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (d["tool"], d["client"], d["phone"], n.strftime("%Y-%m-%d"), due,
             d["days"], d["price"], total, d["deposit"], m.from_user.id),
        )
        rid = cur.lastrowid
    add_record("kirim", "Asbob ijarasi", total, m.text, "Oflayn", d["client"],
               f"{d['tool']} ({d['days']} kun)", m.from_user.id, rental_id=rid)
    await m.answer(
        f"✅ Ijara yozildi #{rid}\n🛠 {d['tool']} → {d['client']}\n"
        f"📅 Qaytarish: {due}\n💰 {fmt(total)} so'm ({m.text}) • garov: {fmt(d['deposit'])}",
        reply_markup=MAIN,
    )


@dp.message()
async def fallback(m: Message):
    await m.answer("Menyudan tanlang 👇", reply_markup=MAIN)


# ====================== AVTOMATIK XABARLAR ======================
async def notify_all(bot: Bot, text: str = None, document=None, caption=None):
    for uid in ALLOWED:
        try:
            if document:
                await bot.send_document(uid, document, caption=caption)
            else:
                await bot.send_message(uid, text, parse_mode="HTML")
        except Exception as e:
            logging.warning("Xabar yuborilmadi (%s): %s", uid, e)


async def scheduler(bot: Bot):
    done = set()
    while True:
        try:
            n = now()
            key = n.strftime("%Y-%m-%d")
            if n.hour == HOUR_REMINDER and (key, "rem") not in done:
                done.add((key, "rem"))
                with conn() as con:
                    rows = con.execute(
                        "SELECT id,tool,client,phone,start_date,due_date,days,price_day,total,deposit FROM rentals WHERE returned=0 AND due_date <= ? ORDER BY due_date",
                        (key,),
                    ).fetchall()
                if rows:
                    await notify_all(bot, "⏰ <b>Bugun qaytarilishi kerak / muddati o'tgan asboblar:</b>\n\n"
                                     + "\n\n".join(rental_text(r) for r in rows))
            if n.hour == HOUR_REPORT and (key, "rep") not in done:
                done.add((key, "rep"))
                await notify_all(bot, build_report(key, key, f"Kunlik hisobot: {n.strftime('%d.%m.%Y')}"))
            if n.hour == HOUR_BACKUP and (key, "bak") not in done:
                done.add((key, "bak"))
                await notify_all(bot, document=FSInputFile(make_backup(), filename=f"zaxira_{key}.db"),
                                 caption="🛟 Kunlik baza nusxasi")
            if len(done) > 50:
                done = {x for x in done if x[0] == key}
        except Exception:
            logging.exception("Scheduler xatosi")
        await asyncio.sleep(60)


async def main():
    if not BOT_TOKEN or not ALLOWED:
        raise SystemExit("BOT_TOKEN va ALLOWED_IDS o'zgaruvchilarini kiriting!")
    init_db()
    bot = Bot(BOT_TOKEN)
    asyncio.create_task(scheduler(bot))
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
